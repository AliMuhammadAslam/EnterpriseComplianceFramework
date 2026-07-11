import os
from utils.logger import logger_instance


class DocumentParser:
    """Parses PDF, DOCX, and TXT files into plain text."""

    def __init__(self):
        self.logger = logger_instance.get_logger("document_parser")

    def parse_file(self, file_path: str) -> str:
        """Parse a file and return its text content.
        
        Supports: .pdf, .docx, .txt
        """
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return self._parse_pdf(file_path)
        elif ext == ".docx":
            return self._parse_docx(file_path)
        elif ext == ".txt":
            return self._parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    def _parse_pdf(self, file_path: str) -> str:
        """Extract text from PDF using PyPDF2."""
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(file_path)
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            text = "\n\n".join(text_parts)
            self.logger.info(
                f"Parsed PDF: {os.path.basename(file_path)} "
                f"({len(reader.pages)} pages, {len(text)} chars)"
            )
            return text
        except Exception as e:
            self.logger.error(f"Error parsing PDF {file_path}: {e}")
            raise

    def _parse_docx(self, file_path: str) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            text = "\n\n".join(text_parts)
            self.logger.info(
                f"Parsed DOCX: {os.path.basename(file_path)} "
                f"({len(doc.paragraphs)} paragraphs, {len(text)} chars)"
            )
            return text
        except Exception as e:
            self.logger.error(f"Error parsing DOCX {file_path}: {e}")
            raise

    def _parse_txt(self, file_path: str) -> str:
        """Read plain text file."""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            self.logger.info(
                f"Parsed TXT: {os.path.basename(file_path)} "
                f"({len(text)} chars)"
            )
            return text
        except UnicodeDecodeError:
            # Fallback to latin-1 encoding
            with open(file_path, "r", encoding="latin-1") as f:
                text = f.read()
            return text
